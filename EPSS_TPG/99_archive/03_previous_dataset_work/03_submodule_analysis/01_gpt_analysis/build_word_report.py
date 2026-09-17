#!/usr/bin/env python3
"""Create editable Word versions of the verified Markdown report."""

from __future__ import annotations

import re
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


BASE = Path(__file__).resolve().parent
REPORT_DIR = BASE / "Final_Explanation_Report"
MD_PATH = REPORT_DIR / "merged_verified_social_media_epss_analysis.md"
DOCX_PATH = REPORT_DIR / "merged_verified_social_media_epss_analysis.docx"
DOC_PATH = REPORT_DIR / "merged_verified_social_media_epss_analysis.doc"


def set_cell_text(cell, value: str, *, bold: bool, font_size: float) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(value.strip())
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(font_size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def keep_table_row_together(row) -> None:
    properties = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    properties.append(cant_split)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.extend([field_begin, instruction, field_end])


def clean_inline(value: str) -> str:
    value = re.sub(r"!\[([^]]*)\]\([^)]+\)", r"\1", value)
    value = re.sub(r"\[([^]]+)\]\(([^)]+)\)", lambda m: m.group(1) if m.group(1) == m.group(2) else f"{m.group(1)} ({m.group(2)})", value)
    value = value.replace("**", "").replace("__", "")
    value = value.replace("`", "")
    return value.strip()


def add_inline_runs(paragraph, text: str) -> None:
    link_pattern = re.compile(r"\[([^]]+)\]\(([^)]+)\)")
    text = link_pattern.sub(lambda m: m.group(1) if m.group(1) == m.group(2) else f"{m.group(1)} ({m.group(2)})", text)
    pattern = re.compile(r"(\*\*.*?\*\*|`.*?`|\*[^*]+?\*)")
    position = 0
    for match in pattern.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position : match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(9)
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def table_rows(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        cells = [clean_inline(cell) for cell in line.strip().strip("|").split("|")]
        rows.append(cells)
    return rows


def is_separator_row(row: list[str]) -> bool:
    return bool(row) and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in row)


def add_markdown_table(document: Document, lines: list[str]) -> None:
    rows = table_rows(lines)
    if len(rows) > 1 and is_separator_row(rows[1]):
        rows.pop(1)
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    table.autofit = True
    compact_size = 7.0 if width >= 7 else 7.5 if width >= 5 else 8.5
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            set_cell_text(
                table.cell(row_index, column_index),
                value,
                bold=row_index == 0,
                font_size=compact_size,
            )
            if row_index == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "D9E6F5")
                table.cell(row_index, column_index)._tc.get_or_add_tcPr().append(shading)
    repeat_table_header(table.rows[0])
    for row in table.rows:
        keep_table_row_together(row)
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, color in [
        ("Title", 21, "111111"),
        ("Heading 1", 15, "1F4E79"),
        ("Heading 2", 12, "1F4E79"),
        ("Heading 3", 10.5, "1F4E79"),
    ]:
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True


def build_docx() -> None:
    document = Document()
    configure_styles(document)
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.9)
    section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.7)
    add_page_number(section.footer.paragraphs[0])

    document.core_properties.title = "Social Media and EPSS Analysis"
    document.core_properties.subject = "Verified social-media, temporal EPSS, similar-CVE comparison, and KEV analysis"

    lines = MD_PATH.read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()
        if not stripped:
            index += 1
            continue

        if stripped.startswith("|"):
            block = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                block.append(lines[index])
                index += 1
            add_markdown_table(document, block)
            continue

        image_match = re.fullmatch(r"!\[([^]]*)\]\(([^)]+)\)", stripped)
        if image_match:
            path = Path(image_match.group(2))
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(12)
            paragraph.paragraph_format.space_after = Pt(6)
            run = paragraph.add_run()
            run.add_picture(str(path), width=Inches(6.15))
            index += 1
            continue

        heading_match = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading_match:
            level = len(heading_match.group(1))
            text = clean_inline(heading_match.group(2))
            paragraph = document.add_paragraph(text, style="Title" if level == 1 else f"Heading {level - 1}")
            if level == 1:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            index += 1
            continue

        if stripped.startswith("> "):
            paragraph = document.add_paragraph(style="Quote")
            add_inline_runs(paragraph, stripped[2:])
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(8)
            index += 1
            continue

        bullet_match = re.match(r"^-\s+(.+)$", stripped)
        numbered_match = re.match(r"^(\d+)\.\s+(.+)$", stripped)
        if bullet_match:
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline_runs(paragraph, bullet_match.group(1))
            paragraph.paragraph_format.space_after = Pt(3)
            index += 1
            continue
        if numbered_match:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.65)
            paragraph.paragraph_format.first_line_indent = Cm(-0.4)
            paragraph.add_run(f"{numbered_match.group(1)}. ")
            add_inline_runs(paragraph, numbered_match.group(2))
            paragraph.paragraph_format.space_after = Pt(3)
            index += 1
            continue

        paragraph = document.add_paragraph()
        add_inline_runs(paragraph, stripped)
        if stripped.startswith("**Figure"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(12)
            paragraph.paragraph_format.keep_with_next = False
        elif stripped.startswith("**Example") or stripped.startswith("**Small example"):
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(8)
            paragraph.paragraph_format.keep_together = True
        elif stripped.startswith("**How to read") or stripped.startswith("**Table notes"):
            paragraph.paragraph_format.space_after = Pt(9)
        elif stripped in {"All posted CVEs with selected comparison CVEs:", "More strictly cleaned first-post group:"}:
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_after = Pt(3)
        index += 1

    document.save(DOCX_PATH)


def convert_to_doc() -> None:
    if DOC_PATH.exists():
        DOC_PATH.unlink()
    command = [
        "libreoffice",
        "-env:UserInstallation=file:///tmp/lo_word_report_profile_escalated",
        "--headless",
        "--convert-to",
        "doc:MS Word 97",
        "--outdir",
        str(REPORT_DIR),
        str(DOCX_PATH),
    ]
    result = subprocess.run(command, check=True, capture_output=True, text=True)
    if not DOC_PATH.exists():
        raise RuntimeError(f"LibreOffice did not create {DOC_PATH}: {result.stdout} {result.stderr}")


def main() -> None:
    build_docx()
    convert_to_doc()
    print(DOCX_PATH)
    print(DOC_PATH)


if __name__ == "__main__":
    main()
